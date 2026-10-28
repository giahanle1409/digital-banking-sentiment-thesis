#!/usr/bin/env python3
"""
PHASE 2a (v3) - Build the stratified Gold-set for manual sentiment labelling
============================================================================
Corpus after Phase 1 v2: ~12,900 reviews (Google Play + App Store, Jan 2024 - Oct 2026).

DESIGN
  * Strata  = App x Platform x Score = 3 x 2 x 5 = 30 cells.
  * Tempered allocation  n_cell  proportional to  N_cell ** ALPHA   (ALPHA = 0.5 -> square root of the cell size),
    with a floor of MIN_PER_CELL = 25 reviews per cell (all of them if the cell is smaller). The allocation is solved
    by water-filling so the total hits N exactly and no cell exceeds its population.
  * Sampling weights  w = N_h / n_h  are saved in gold_key.csv. The gold set is NOT representative of the corpus:
    estimate corpus prevalence with weighted sums (sum(w * 1[label=k]) / sum(w)).
  * Optional --cand-share q (0..1): inside every cell a share q of the sample is drawn from "Neutral/Mixed candidates"
    (reviews with a contrast cue such as 'nhưng', 'tuy nhiên', or a question / suggestion cue), the rest from the other
    reviews. Each (cell x candidate flag) is a sub-stratum with its own weight N_h/n_h, so weighted estimates stay valid.
  * Blinding: annotator sheets contain gold_id, Platform and the review text only. No star score, no app name,
    no thumbs-up count. The key (with stars and weights) stays in gold_key.csv - never send it to a coder.

WORKFLOW (see the thesis plan in the chat message)
  pilot (100, both coders)  ->  Cohen's kappa  ->  refine + FREEZE guideline  ->  Coder 1 labels the remaining
  1,900  ->  Coder 2 labels 300 more (400 double-coded in total, pilot included)  ->  02c_merge_annotations.py.

OUTPUTS (data/gold/)
  gold_key.csv                       private key: strata, Score, weight, flags, phase
  gold_pilot_coder1.xlsx / _coder2.xlsx      the 100 pilot reviews (same rows, different order)
  gold_main_coder1.xlsx              the remaining N-100 reviews for Coder 1
  gold_main_coder2_extra.xlsx        the other double-coded reviews (double_n - 100) for Coder 2
  Every sheet has EXACTLY these columns: gold_id | Platform | text_raw | label | lean_if_mixed | notes
  with dropdowns on label (Positive, Negative, Neutral, Mixed, Junk/Spam) and lean_if_mixed (Positive, Negative, Neutral);
  rows are coloured red when label = Mixed but lean_if_mixed is empty, and orange when lean_if_mixed is filled
  although label is not Mixed.
  gold_design_report.csv             N_cell, n_cell, weight per cell (Chapter 3 / appendix table)
  gold_prelabels_from_previous_round.csv   (only with --reuse-key/--reuse-labels) labels already given in the
                                     previous round to identical review texts - verify, do not copy blindly

USAGE
  python 02a_make_goldset_v3.py --clean data/reviews_clean.csv                      # 2,000 reviews, pure sqrt design
  python 02a_make_goldset_v3.py --clean data/reviews_clean.csv --n 2700 --cand-share 0.25   # scale-up + oversampling
"""
import argparse
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

LABELS = "Positive,Negative,Neutral,Mixed,Junk/Spam"
LEANS = "Positive,Negative,Neutral"
GUIDELINE = [
    ("Unit of annotation", "Label the OVERALL attitude the writer expresses toward the app, the bank or the service, using ONLY the text. "
                           "The star rating is hidden on purpose; never infer the label from it."),
    ("Decision order", "1) Junk/Spam? 2) Is any evaluation present? (if not -> Neutral) 3) Are BOTH positive and negative evaluations of "
                       "service attributes present? (-> Mixed) 4) Otherwise Positive or Negative."),
    ("Positive", "Overall approval: praise, satisfaction, thanks, recommendation ('tốt', 'mượt', 'rất hài lòng')."),
    ("Negative", "Overall disapproval: complaint, anger, warning to others, accusation of fraud, a reported failure, an insult ('tệ', 'lừa đảo')."),
    ("Neutral", "No clear evaluation (a question, a factual report, a suggestion without emotion) or a balanced statement."),
    ("Mixed", "Positive AND negative evaluations of service attributes in the same review ('app đẹp nhưng lỗi liên tục')."),
    ("lean_if_mixed", "REQUIRED when label = Mixed, otherwise leave empty: the polarity the review leans toward overall "
                      "(Positive, Negative, or Neutral if truly balanced). It is used for the 3-class fallback label."),
    ("Junk/Spam", "Not an evaluation: referral codes, promotional text, gibberish, off-topic. Excluded from training and evaluation."),
    ("Sarcasm", "Label the intended meaning ('tuyệt vời, chuyển tiền 3 ngày chưa tới' -> Negative). Add a note."),
    ("Negation", "'không tệ' leans Positive; 'không tốt' is Negative. Read the whole sentence."),
    ("Short generic text", "'ok', 'tốt' -> Positive; 'tệ', 'lừa đảo' -> Negative. Decide once, write it here, then freeze."),
    ("Teencode / unaccented", "Interpret normally ('ko dc' = không được). Do not skip such rows."),
    ("Developer replies / quotes", "Ignore pasted bank replies; label the customer's own voice."),
    ("Process", "Pilot of 100 reviews with both coders -> Cohen's kappa -> refine and FREEZE this guideline -> label the rest."),
]
# Cues for Neutral / Mixed candidates (used only as a sampling device; labels are always assigned by the coders)
MIX_CUE = re.compile(r"\b(?:nhưng|tuy nhiên|tuy|nhưng mà|mặc dù|có điều|chỉ tiếc|tiếc là|ngoại trừ)\b", re.I)
NEU_CUE = re.compile(r"(?:làm sao|như thế nào|bao giờ|khi nào|cho (?:mình )?hỏi|hướng dẫn|xin hỏi|có cách nào|đề xuất|đề nghị|"
                     r"mong (?:app|ngân hàng|bên)|nên (?:thêm|bổ sung|có)|\?)", re.I)


def allocate(sizes: pd.Series, n_total: int, alpha: float, floor: int) -> pd.Series:
    """Tempered allocation with floor and cap (water-filling: cells that hit their cap release budget)."""
    cap = sizes.values.astype(float)
    n = np.minimum(cap, floor)
    remaining = n_total - n.sum()
    w = cap ** alpha
    for _ in range(200):
        room = cap - n
        active = room > 1e-9
        if remaining < 0.5 or not active.any():
            break
        add = np.minimum(remaining * w * active / (w * active).sum(), room)
        n += add
        remaining -= add.sum()
    out = np.floor(n).astype(int)
    for i in np.argsort(-(n - out)):               # hand out the rounding leftovers
        if out.sum() >= n_total:
            break
        if out[i] < cap[i]:
            out[i] += 1
    return pd.Series(out, index=sizes.index)


def write_sheet(df: pd.DataFrame, path: Path) -> None:
    """Annotator sheet: gold_id | Platform | text_raw | label | lean_if_mixed | notes (no score, no app, no thumbs-up)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "annotate"
    ws.append(["gold_id", "Platform", "text_raw", "label", "lean_if_mixed", "notes"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    for _, r in df.iterrows():
        ws.append([r["gold_id"], r["Platform"], str(r["text_raw"]), None, None, None])
        ws.cell(row=ws.max_row, column=3).data_type = "s"                  # never let Excel read a review as a formula
    last = len(df) + 1
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 9, 12
    ws.column_dimensions["C"].width, ws.column_dimensions["D"].width = 100, 12
    ws.column_dimensions["E"].width, ws.column_dimensions["F"].width = 15, 30
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    dv_label = DataValidation(type="list", formula1=f'"{LABELS}"', allow_blank=True, showErrorMessage=True,
                              errorTitle="Invalid label", error="Choose a value from the list.")
    dv_lean = DataValidation(type="list", formula1=f'"{LEANS}"', allow_blank=True, showErrorMessage=True,
                             errorTitle="Invalid lean", error="Choose Positive, Negative or Neutral.")
    ws.add_data_validation(dv_label); ws.add_data_validation(dv_lean)
    dv_label.add(f"D2:D{last}"); dv_lean.add(f"E2:E{last}")
    # visual checks: red = Mixed without lean; orange = lean given but label is not Mixed
    ws.conditional_formatting.add(f"A2:F{last}", FormulaRule(formula=['AND($D2="Mixed",$E2="")'],
                                                             fill=PatternFill("solid", bgColor="FFC7CE")))
    ws.conditional_formatting.add(f"A2:F{last}", FormulaRule(formula=['AND($D2<>"Mixed",$E2<>"")'],
                                                             fill=PatternFill("solid", bgColor="FFD966")))
    ws.freeze_panes = "D2"
    g = wb.create_sheet("guideline")
    g.append(["Topic", "Rule"])
    for t, r in GUIDELINE:
        g.append([t, r])
    g.column_dimensions["A"].width, g.column_dimensions["B"].width = 26, 140
    for row in g.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def proportional_by(df: pd.DataFrame, col: str, k: int, rng, pool=None) -> pd.Index:
    """Pick k rows of `pool` (default df) stratified (proportionally) by `col`."""
    pool = df if pool is None else pool
    share = pool[col].value_counts(normalize=True)
    idx = []
    for v, s in share.items():
        sub = pool[pool[col] == v]
        idx.extend(sub.sample(n=min(len(sub), int(round(k * s))), random_state=rng).index)
    idx = list(dict.fromkeys(idx))
    if len(idx) > k:
        idx = list(rng.permutation(idx)[:k])
    if len(idx) < k:                                                   # rounding: top up so exactly k are chosen
        left = pool.drop(idx)
        idx += list(left.sample(n=k - len(idx), random_state=rng).index)
    return pd.Index(idx)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--clean", default="data/reviews_clean.csv")
    ap.add_argument("--outdir", default="data/gold")
    ap.add_argument("--n", type=int, default=2000, help="gold-set size (2,000 base; 2,700 recommended scale-up)")
    ap.add_argument("--alpha", type=float, default=0.5, help="0.5 = allocation proportional to sqrt(cell size)")
    ap.add_argument("--min-per-cell", type=int, default=25)
    ap.add_argument("--pilot-n", type=int, default=100, help="pilot reviews labelled by both coders")
    ap.add_argument("--double-n", type=int, default=400, help="total double-coded reviews (pilot included)")
    ap.add_argument("--cand-share", type=float, default=0.0, help="share of each cell drawn from Neutral/Mixed candidates")
    ap.add_argument("--reuse-key", default=None, help="previous gold_key.csv (needs text_raw)")
    ap.add_argument("--reuse-labels", default=None, help="previous gold_labeled_full.csv (needs gold_id, label5)")
    ap.add_argument("--seed", type=int, default=42)
    a = ap.parse_args()

    out = Path(a.outdir); out.mkdir(parents=True, exist_ok=True)
    df = pd.read_csv(a.clean)
    df["cell"] = df["App"] + "|" + df["Platform"] + "|" + df["Score"].astype(str)
    df["cand"] = (df["text_raw"].str.contains(MIX_CUE) & (df["n_words_clean"] >= 6)) | df["text_raw"].str.contains(NEU_CUE)

    sizes = df.groupby("cell").size()
    alloc = allocate(sizes, a.n, a.alpha, a.min_per_cell)
    rng = np.random.RandomState(a.seed)

    parts = []
    for cell, k in alloc.items():
        pool = df[df["cell"] == cell]
        if a.cand_share <= 0:                                          # pure design: one stratum per cell
            s = pool.sample(n=int(k), random_state=rng).copy()
            s["N_h"], s["n_h"] = len(pool), int(k)
            parts.append(s)
            continue
        k_c = min(int((pool["cand"]).sum()), int(round(a.cand_share * k)))
        k_o = k - k_c
        other = pool[~pool["cand"]]
        if k_o > len(other):                                           # not enough non-candidates: top up with candidates
            k_c += k_o - len(other); k_o = len(other)
        for flag, kk in ((True, k_c), (False, k_o)):
            sub = pool[pool["cand"] == flag]
            if kk > 0:
                s = sub.sample(n=kk, random_state=rng).copy()
                s["N_h"], s["n_h"] = len(sub), kk
                parts.append(s)
    gold = pd.concat(parts)
    gold["weight"] = gold["N_h"] / gold["n_h"]                          # sampling weight = N_h / n_h
    gold = gold.sample(frac=1, random_state=a.seed).reset_index(drop=True)       # shuffle: no ordering by app/score
    gold.insert(0, "gold_id", [f"G{i:04d}" for i in range(1, len(gold) + 1)])

    # double coding: pilot is a subset of the double-coded set; both stratified by star score
    pil_idx = proportional_by(gold, "Score", a.pilot_n, rng)
    rest = gold.drop(pil_idx)
    ext_idx = proportional_by(gold, "Score", a.double_n - len(pil_idx), rng, pool=rest)
    gold["phase"] = "main"
    gold.loc[pil_idx, "phase"] = "pilot"
    gold["double_code"] = gold.index.isin(pil_idx) | gold.index.isin(ext_idx)

    key_cols = ["gold_id", "review_id", "App", "Platform", "Score", "cell", "cand", "N_h", "n_h", "weight",
                "phase", "double_code", "n_words_clean", "text_raw"]
    gold[key_cols].to_csv(out / "gold_key.csv", index=False, encoding="utf-8-sig")

    pilot = gold[gold.phase == "pilot"]
    write_sheet(pilot, out / "gold_pilot_coder1.xlsx")
    write_sheet(pilot.sample(frac=1, random_state=a.seed + 1), out / "gold_pilot_coder2.xlsx")
    write_sheet(gold[gold.phase == "main"], out / "gold_main_coder1.xlsx")
    extra = gold[(gold.phase == "main") & gold.double_code].sample(frac=1, random_state=a.seed + 2)
    write_sheet(extra, out / "gold_main_coder2_extra.xlsx")

    rep = (pd.DataFrame({"N_cell": sizes, "n_gold": gold.groupby("cell").size()}).fillna(0).astype(int)
           .assign(weight=lambda t: t["N_cell"] / t["n_gold"].replace(0, np.nan)).reset_index())
    rep[["App", "Platform", "Score"]] = rep["cell"].str.split("|", expand=True)
    rep.drop(columns="cell").to_csv(out / "gold_design_report.csv", index=False)

    if a.reuse_key and a.reuse_labels:
        ok = pd.read_csv(a.reuse_key)[["gold_id", "App", "Platform", "text_raw"]]
        ol = pd.read_csv(a.reuse_labels)[["gold_id", "label5"]]
        old = ok.merge(ol, on="gold_id").rename(columns={"gold_id": "old_gold_id"}).drop_duplicates(["App", "Platform", "text_raw"])
        pre = gold[["gold_id", "App", "Platform", "text_raw"]].merge(old, on=["App", "Platform", "text_raw"])
        pre[["gold_id", "old_gold_id", "label5"]].to_csv(out / "gold_prelabels_from_previous_round.csv", index=False)
        print(f"Reusable labels from the previous round: {len(pre)} of {len(gold)} reviews")

    # ---------------- report ----------------
    print(f"Gold-set: {len(gold)} | pilot: {int((gold.phase == 'pilot').sum())} | double-coded: {int(gold.double_code.sum())}"
          f" | candidates: {int(gold.cand.sum())}")
    print(rep.pivot_table(index=["App", "Platform"], columns="Score", values="n_gold", aggfunc="sum").to_string())
    print("weights: min %.2f  median %.2f  max %.2f" % (gold.weight.min(), gold.weight.median(), gold.weight.max()))
    print("Share of gold with <4 words: %.3f | corpus: %.3f" % ((gold.n_words_clean < 4).mean(), (df.n_words_clean < 4).mean()))
    # expected yield of rare classes from the previous round (2,000 reviews: Neutral 2.6%, Mixed 1.4%; among candidates
    # Neutral 3.6%, Mixed 6.8%; among non-candidates Neutral ~2.4%, Mixed ~0.7%)
    nc, no = int(gold.cand.sum()), int((~gold.cand).sum())
    for share in (0.15,):
        print("Expected Neutral/Mixed in the whole gold set: %.0f / %.0f ; in a %d%% test split: %.0f / %.0f" % (
            nc * .036 + no * .024, nc * .068 + no * .007, share * 100,
            share * (nc * .036 + no * .024), share * (nc * .068 + no * .007)))


if __name__ == "__main__":
    main()
